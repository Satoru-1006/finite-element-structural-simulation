from pathlib import Path
from textwrap import dedent

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\86198\Desktop\task8_nature_output")
FIG = OUT / "figures"
SRC_MODEL = OUT / "source_model.m"
ORIG_FIG = OUT / "original_figures"
FINAL = OUT / "Task8_Nature_final_two_column_layout_v5_formula_clean.docx"
FINAL_ZH = OUT / "任务8_Nature双栏公式修正版_v5.docx"


def rfonts(run, latin="Times New Roman", east_asia="宋体"):
    run.font.name = latin
    run._element.rPr.rFonts.set(qn("w:eastAsia"), east_asia)


def set_columns(section, num=1, space=720):
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(space))


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_borders(table, color="BDBDBD", size="4"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)
        borders.append(tag)
    tbl_pr.append(borders)


def set_table_width(table, width_cm):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.xpath("./w:tblW")
    if tbl_w:
        tbl_w = tbl_w[0]
    else:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(int(width_cm * 567)))
    tbl_w.set(qn("w:type"), "dxa")


def keep_paragraph(p, keep_next=False, keep_lines=True):
    p_pr = p._p.get_or_add_pPr()
    if keep_next:
        p_pr.append(OxmlElement("w:keepNext"))
    if keep_lines:
        p_pr.append(OxmlElement("w:keepLines"))


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录（在 Word 中右键选择“更新域”即可刷新页码）"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)
    rfonts(run, "Times New Roman", "宋体")


def add_para(doc, text="", style=None, first_indent=True):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    if first_indent and style is None:
        p.paragraph_format.first_line_indent = Cm(0.74)
    run = p.add_run(text)
    rfonts(run)
    run.font.size = Pt(10.5)
    return p


def add_run_para(doc, parts, first_indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    if first_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    for text, italic in parts:
        run = p.add_run(text)
        rfonts(run)
        run.font.size = Pt(10.5)
        run.italic = italic
    return p


def add_caption(doc, label, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.first_line_indent = Cm(0)
    r1 = p.add_run(label)
    r1.bold = True
    rfonts(r1, "Arial", "等线")
    r1.font.size = Pt(9.5)
    r2 = p.add_run(" " + text)
    rfonts(r2, "Arial", "等线")
    r2.font.size = Pt(9.5)
    keep_paragraph(p, keep_lines=True)


def add_table_title(doc, label, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.first_line_indent = Cm(0)
    keep_paragraph(p, keep_next=True)
    r1 = p.add_run(label)
    r1.bold = True
    rfonts(r1, "Arial", "等线")
    r1.font.size = Pt(9.5)
    r2 = p.add_run(" " + title)
    rfonts(r2, "Arial", "等线")
    r2.font.size = Pt(9.5)


def table_cell(cell, text, bold=False, size=8.6, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.first_line_indent = Cm(0)
    run = p.add_run(text)
    run.bold = bold
    rfonts(run, "Arial", "等线")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    if shade:
        set_cell_shading(cell, shade)


def add_compact_table(doc, headers, rows, widths_cm):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "C8C8C8", "3")
    for i, (cell, text) in enumerate(zip(table.rows[0].cells, headers)):
        table_cell(cell, text, bold=True, shade="EFEFEF")
        cell.width = Cm(widths_cm[i])
    cant_split(table.rows[0])
    for row in rows:
        cells = table.add_row().cells
        for i, (cell, text) in enumerate(zip(cells, row)):
            table_cell(cell, text)
            cell.width = Cm(widths_cm[i])
        cant_split(table.rows[-1])
    return table


def add_omml_equation(doc, omml, number, wide=True):
    if wide:
        sec = doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(sec, 1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "FFFFFF", "0")
    total_width = 16.4 if wide else 7.45
    eq_width = 15.2 if wide else 6.55
    no_width = 1.0 if wide else 0.85
    set_table_width(table, total_width)
    table.columns[0].width = Cm(eq_width)
    table.columns[1].width = Cm(no_width)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    right.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.0
    p._p.append(parse_xml(omml))
    rp = right.paragraphs[0]
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rr = rp.add_run(f"({number})")
    rfonts(rr, "Times New Roman", "宋体")
    rr.font.size = Pt(10)
    if wide:
        sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
        set_columns(sec2, 2)
    return table


def omath_text(text):
    return f'<m:oMathPara {nsdecls("m")}><m:oMath><m:r><m:t>{text}</m:t></m:r></m:oMath></m:oMathPara>'


def mf(num, den):
    return f"<m:f><m:num><m:r><m:t>{num}</m:t></m:r></m:num><m:den><m:r><m:t>{den}</m:t></m:r></m:den></m:f>"


def omath_parts(parts):
    body = []
    for item in parts:
        if isinstance(item, tuple):
            body.append(mf(item[0], item[1]))
        else:
            body.append(f"<m:r><m:t>{item}</m:t></m:r>")
    return f'<m:oMathPara {nsdecls("m")}><m:oMath>{"".join(body)}</m:oMath></m:oMathPara>'


def omml_matrix(rows, beg="[", end="]"):
    matrix_rows = []
    for row in rows:
        cells = "".join(f"<m:e><m:r><m:t>{cell}</m:t></m:r></m:e>" for cell in row)
        matrix_rows.append(f"<m:mr>{cells}</m:mr>")
    matrix = f"<m:m>{''.join(matrix_rows)}</m:m>"
    return (
        f'<m:d><m:dPr><m:begChr m:val="{beg}"/><m:endChr m:val="{end}"/></m:dPr>'
        f"<m:e>{matrix}</m:e></m:d>"
    )


def omath_matrix_constitutive():
    return f"""
    <m:oMathPara {nsdecls("m")}><m:oMath>
      {omml_matrix([["εₓ"], ["εᵧ"], ["γₓᵧ"]])}
      <m:r><m:t>=</m:t></m:r>
      {mf("1", "E")}
      {omml_matrix([["1", "-ν", "0"], ["-ν", "1", "0"], ["0", "0", "2(1+ν)"]])}
      {omml_matrix([["σₓ"], ["σᵧ"], ["τₓᵧ"]])}
    </m:oMath></m:oMathPara>
    """


def add_omml_equation_group(doc, equations):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 1)
    table = doc.add_table(rows=len(equations), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "FFFFFF", "0")
    for row, (omml, number) in zip(table.rows, equations):
        cant_split(row)
        eq_cell = row.cells[0]
        no_cell = row.cells[1]
        eq_cell.width = Cm(15.2)
        no_cell.width = Cm(1.0)
        p = eq_cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p._p.append(parse_xml(omml))
        np = no_cell.paragraphs[0]
        np.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        np.paragraph_format.space_before = Pt(0)
        np.paragraph_format.space_after = Pt(0)
        nr = np.add_run(f"({number})")
        rfonts(nr, "Times New Roman", "宋体")
        nr.font.size = Pt(10)
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec2, 2)
    return table


def add_equation_para(doc, omml, number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(0 if number else 2)
    p.paragraph_format.line_spacing = 1.0
    p._p.append(parse_xml(omml))
    if number is not None:
        pn = doc.add_paragraph()
        pn.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pn.paragraph_format.first_line_indent = Cm(0)
        pn.paragraph_format.space_before = Pt(0)
        pn.paragraph_format.space_after = Pt(2)
        rn = pn.add_run(f"({number})")
        rfonts(rn, "Times New Roman", "宋体")
        rn.font.size = Pt(9.5)
    return p


def add_wide_figure(doc, filename, caption_label, caption_text, width_cm=15.6):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(FIG / filename), width=Cm(width_cm))
    add_caption(doc, caption_label, caption_text)
    sec2 = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec2, 2)


def add_appendix_code(doc, code):
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 1)
    doc.add_heading("附录 A  MATLAB 核心代码", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("以下代码保留原 MATLAB 语法，其中 nu^2、n^2 等写法属于代码表达，不按正文公式格式改写。")
    rfonts(r, "Times New Roman", "宋体")
    r.font.size = Pt(10)
    lines = code.splitlines()
    selected = []
    keep = False
    for line in lines:
        if "%% 基本参数" in line:
            keep = True
        if keep:
            selected.append(line)
        if "u2 = u2 - u2(:,1);" in line:
            break
    box = doc.add_table(rows=1, cols=1)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(box, "D0D0D0", "4")
    cell = box.cell(0, 0)
    set_cell_shading(cell, "F5F5F5")
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.first_line_indent = Cm(0)
    for i, line in enumerate(selected):
        if i:
            p.add_run("\n")
        run = p.add_run(line)
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(8.5)


def add_original_figures_appendix(doc):
    if not ORIG_FIG.exists():
        return
    images = sorted(ORIG_FIG.glob("original_figure_*.png"))
    if not images:
        return
    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_columns(sec, 1)
    doc.add_heading("附录 B  原始 MATLAB 图 1-17", level=1)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run("以下图片为原始课程报告中的 MATLAB 输出图，作为原始结果记录保留。正文中的 Figure 1-4 已将这些结果重新整理为期刊风格多面板图。")
    rfonts(r, "Times New Roman", "宋体")
    r.font.size = Pt(9.5)

    table = doc.add_table(rows=0, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table, "FFFFFF", "0")
    for start in range(0, len(images), 3):
        row = table.add_row()
        cant_split(row)
        for j in range(3):
            cell = row.cells[j]
            cell.width = Cm(5.45)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            cell.text = ""
            idx = start + j
            if idx >= len(images):
                continue
            img = images[idx]
            p_img = cell.paragraphs[0]
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(1)
            p_img.paragraph_format.space_after = Pt(1)
            p_img.add_run().add_picture(str(img), width=Cm(5.05))
            p_cap = cell.add_paragraph()
            p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_cap.paragraph_format.space_after = Pt(5)
            rr = p_cap.add_run(f"原始图 {idx + 1}")
            rr.bold = True
            rfonts(rr, "Arial", "等线")
            rr.font.size = Pt(8.5)


doc = Document()
sec = doc.sections[0]
sec.orientation = WD_ORIENT.PORTRAIT
sec.page_width = Cm(21)
sec.page_height = Cm(29.7)
sec.top_margin = Cm(1.8)
sec.bottom_margin = Cm(1.8)
sec.left_margin = Cm(1.7)
sec.right_margin = Cm(1.7)
set_columns(sec, 1)

styles = doc.styles
styles["Normal"].font.name = "Times New Roman"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"].font.size = Pt(10.5)
for name, size in [("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 11)]:
    style = styles[name]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "等线")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(8 if name == "Heading 1" else 5)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.keep_with_next = True

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(5)
r = title.add_run("矩形薄板单向拉伸的平面弹性建模与端部扰动分析")
r.bold = True
rfonts(r, "Arial", "等线")
r.font.size = Pt(18)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("有限元分析课程任务 8  Nature-style 课程报告")
rfonts(r, "Arial", "等线")
r.font.size = Pt(11)
r.italic = True

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = meta.add_run("几何参数：L = 1.0 m，H = 0.4 m，t = 0.01 m；材料参数：E = 210 GPa，ν = 0.30；载荷：q = 100 MPa")
rfonts(r, "Times New Roman", "宋体")
r.font.size = Pt(9.5)

doc.add_heading("摘要", level=1)
add_para(doc, "本报告以右端均布拉伸、左端约束的矩形薄板为对象，建立平面应力条件下的二维弹性边值问题，并以艾力应力函数解释均匀拉伸解析解与完全固定左端级数近似解之间的差异。均匀拉伸解给出 σx = 100 MPa、σy = τxy = 0，以及沿 x 方向线性增长的位移 u 和由泊松效应控制的横向位移 v。相比之下，完全固定左端破坏自由泊松收缩，在固定端附近引入 σy、τxy 与 γxy 的局部二维扰动，使最大 σx 升至约 128.55 MPa；该扰动沿板长方向迅速衰减，远离固定端后恢复到均匀拉伸基准。本分析展示了边界条件对连续体应力场的控制作用，并为后续有限元弱形式离散和结果解释提供理论参照。")
kw = add_para(doc, "关键词：平面应力；艾力应力函数；矩形薄板；端部扰动；有限元分析；Nature-style 科研图", first_indent=False)
kw.runs[0].bold = True

doc.add_heading("目录", level=1)
add_toc(doc)

sec = doc.add_section(WD_SECTION.CONTINUOUS)
set_columns(sec, 2)

doc.add_heading("1 引言", level=1)
add_para(doc, "杆、梁等一维结构问题通常可以通过少量广义位移和常微分方程描述，而二维弹性连续体中的位移、应变和应力均为空间场变量。有限元分析的关键步骤，是先将工程对象转化为可验证的连续体边值问题，再通过弱形式或能量原理完成离散。")
add_para(doc, "本任务选择矩形薄板单向拉伸作为过渡模型。该问题既包含清晰的均匀拉伸解析基准，也包含由完全固定边界引发的端部二维扰动。因此，它适合用于说明边界条件、泊松效应、应力重分布和有限元建模之间的关系。")
add_para(doc, "从课程训练角度看，本问题的价值在于把“结构受拉”这一看似简单的工程情景拆解为三个层次：第一，明确几何区域、材料假设和边界条件；第二，写出平面弹性强形式并识别相容性要求；第三，将解析或级数结果转化为可视化证据，并据此判断有限元离散结果是否合理。这样的分析顺序能够避免直接堆叠云图而缺少力学解释。")
add_para(doc, "本文按照科研论文式结构组织报告。首先给出问题定义和理论方程，随后建立均匀拉伸解析基准与固定端扰动近似，最后用多面板静态图展示位移、应力、应变和扰动衰减。正文中的图、表和公式均按连续编号排版，以便结果分析能够形成“文字解释、图像证据、简短结论”的闭环。")

add_wide_figure(
    doc,
    "Figure_1_problem_setup.png",
    "图 1",
    "矩形薄板几何模型与边界条件示意图。A，几何尺寸与坐标系；B，左端完全固定、右端均布拉力以及上下自由边界；C，平面应力假设下的面内应力分量；D，从工程问题到控制方程、艾力应力函数和场变量输出的分析流程。"
)

doc.add_heading("2 理论基础与边值问题", level=1)
doc.add_heading("2.1 平面应力假设", level=2)
add_para(doc, "薄板厚度 t 远小于平面尺寸 L 和 H，厚度方向约束较弱，因此可取 σz ≈ 0，并采用平面应力模型。基本未知量为位移场 u(x,y) 与 v(x,y)，应力分量为 σx、σy 和 τxy，应变分量为 εx、εy 和 γxy。")
add_para(doc, "平面应力并不意味着结构没有厚度，而是指厚度方向正应力相对于面内分量可以忽略。对于本题，右端拉应力沿板面方向施加，上下边界为自由边界，且厚度仅为 0.01 m，因此采用平面应力模型比平面应变模型更符合薄板受拉的物理图像。")

add_table_title(doc, "表 1", "几何、材料与载荷参数。")
add_compact_table(
    doc,
    ["符号", "取值", "含义"],
    [
        ("L", "1.0 m", "板长"),
        ("H", "0.4 m", "板高"),
        ("t", "0.01 m", "板厚"),
        ("E", "210 GPa", "弹性模量"),
        ("ν", "0.30", "泊松比"),
        ("q", "100 MPa", "右端均布拉应力"),
        ("G", "E/[2(1 + ν)]", "剪切模量"),
    ],
    [1.2, 2.0, 4.2],
)

doc.add_heading("2.2 控制方程", level=2)
add_para(doc, "忽略体力时，二维平面弹性强形式由平衡方程、几何方程和本构方程组成。为避免公式像代码文本，下面的主要方程均以 Word OfficeMath 形式给出，可在 Word 中直接编辑。")
add_equation_para(doc, omath_parts([("∂σₓ", "∂x"), "+", ("∂τₓᵧ", "∂y"), "=0"]))
add_equation_para(doc, omath_parts([("∂τₓᵧ", "∂x"), "+", ("∂σᵧ", "∂y"), "=0"]), 1)
add_equation_para(doc, omath_parts(["εₓ=", ("∂u", "∂x")]))
add_equation_para(doc, omath_parts(["εᵧ=", ("∂v", "∂y")]))
add_equation_para(doc, omath_parts(["γₓᵧ=", ("∂u", "∂y"), "+", ("∂v", "∂x")]), 2)
add_equation_para(doc, omath_matrix_constitutive(), 3)

add_table_title(doc, "表 2", "控制方程与边界条件的紧凑整理。")
add_compact_table(
    doc,
    ["类别", "表达", "物理意义"],
    [
        ("区域 Ω", "0 < x < L, −H/2 < y < H/2", "矩形薄板中面区域"),
        ("平衡方程", "式（1）", "无体力条件下的面内力平衡"),
        ("几何方程", "式（2）", "小变形应变-位移关系"),
        ("本构关系", "式（3）", "各向同性平面应力模型"),
        ("左端固定", "x = 0: u = 0, v = 0", "完全固定边界"),
        ("右端载荷", "x = L: σx = q, τxy = 0", "均布拉应力"),
        ("上下自由", "y = ±H/2: σy = 0, τxy = 0", "自由边界"),
    ],
    [1.45, 3.45, 2.55],
)

doc.add_heading("2.3 边界条件与有限元弱形式的联系", level=2)
add_para(doc, "强形式方程要求位移场具有足够高的可微性，并在每一点上满足平衡关系。有限元方法通常不直接求强形式精确解，而是将平衡方程乘以虚位移并在区域上积分，通过分部积分把高阶导数转移到试函数上。这样可以自然引入力边界条件，并允许用分片多项式近似位移场。")
add_para(doc, "在本题中，左端固定属于位移边界，右端均布拉力和上下自由边界属于力边界。有限元离散时，位移边界通常通过约束节点自由度施加，力边界则通过等效节点力或边界积分施加。若左端同时固定 u 和 v，就会抑制该边界处的横向泊松收缩，从而在固定端附近产生局部二维应力扰动。")

doc.add_heading("3 艾力应力函数与解析基准", level=1)
add_para(doc, "艾力应力函数 Φ 通过二阶偏导生成应力分量，使无体力平衡方程自动满足。其核心作用是把多分量应力求解转化为单个标量函数的构造问题。")
add_equation_para(doc, omath_parts(["σₓ=", ("∂²Φ", "∂y²"), ",   σᵧ=", ("∂²Φ", "∂x²")]))
add_equation_para(doc, omath_parts(["τₓᵧ=-", ("∂²Φ", "∂x∂y")]), 4)
add_equation_para(doc, omath_text("∇⁴Φ=0"), 5)
add_para(doc, "在均匀拉伸条件下，可取 Φ0 = qy²/2，从而得到 σx = q、σy = 0 和 τxy = 0。相应位移场为 u = qx/E，v = −νqy/E。完全固定左端时，边界 v = 0 与自由泊松收缩不相容，必须引入随 x 衰减的端部扰动项。")
add_para(doc, "需要强调的是，本文使用的固定端级数近似主要用于说明扰动的空间分布和衰减规律，而不是宣称给出严格闭式解析解。级数项中的指数因子控制沿 x 方向的衰减，三角函数项控制沿高度方向的正负分布。有限项截断后，结果可作为有限元云图和路径曲线的物理参照。")
add_para(doc, "在后续有限元计算中，均匀拉伸解可用于检验远离约束边界的平均应力水平，固定端级数近似则可用于判断端部附近是否出现合理的应力扰动。若数值云图在远场仍存在明显 σy 或 τxy，通常需要检查边界施加方式、网格质量或载荷等效方式。")

doc.add_heading("4 结果与分析", level=1)
doc.add_heading("4.1 均匀拉伸解", level=2)
add_para(doc, "均匀拉伸解提供了全场响应的解析基准。如图 2 所示，u 随 x 线性增大，v 随 y 呈反对称分布，σx 在整个板内保持 100 MPa，εx 和 εy 均为空间常量。这一结果说明，只要边界允许自由泊松收缩，矩形薄板可以维持近似一维的应力状态。")
add_para(doc, "从数值量级看，q/E = 4.76 × 10⁻⁴，因此板右端的最大轴向位移约为 4.76 × 10⁻⁴ m。横向应变为 −νq/E = −1.43 × 10⁻⁴，说明横向收缩幅值小于纵向拉伸幅值，但其边界相容性对固定端应力分布具有决定性影响。")
add_wide_figure(
    doc,
    "Figure_2_uniform_tension_solution.png",
    "图 2",
    "均匀拉伸解析解。A，x 方向位移场 u；B，y 方向位移场 v；C，σx 应力场；D，εx 与 εy 应变场。图中坐标单位为 m，应力单位为 MPa。"
)

doc.add_heading("4.2 完全固定边界的端部扰动", level=2)
add_para(doc, "完全固定左端抑制了横向自由收缩，使固定端附近出现二维应力重分布。如图 3 所示，σy、τxy 和 γxy 的非零区域主要集中在 x = 0 附近，并随距离固定端增大而快速衰减。这一局部扰动是约束条件与泊松效应不相容的直接结果。")
add_para(doc, "与均匀拉伸解相比，固定端响应的关键变化并不只体现在 u = 0 的位移约束上，更体现在 v = 0 对泊松收缩的抑制。由于材料仍试图发生横向收缩，而边界不允许其自由运动，固定端附近必须通过横向正应力和剪应力重新分配来满足平衡与协调。")
add_para(doc, "图 3 中 σy 和 τxy 采用以零为中心的发散色标，是为了突出扰动的正负性质，而不是夸大其绝对量级。可以看到，扰动区域与左端约束相邻，离开约束后很快减弱；这与 Saint-Venant 原理所描述的局部边界效应相一致。")
add_wide_figure(
    doc,
    "Figure_3_fixed_boundary_solution.png",
    "图 3",
    "左端完全固定条件下的位移、应力与剪应变场。A 和 B，位移分量 u 与 v；C，σx；D，σy；E，τxy；F，γxy。σy、τxy 和 γxy 使用以零为中心的发散色标，灰色阴影表示固定端扰动区。"
)

doc.add_heading("4.3 边界条件对峰值应力的影响", level=2)
add_para(doc, "沿中线 y = 0 的应力曲线进一步量化了端部扰动的衰减过程。如图 4A 所示，完全固定解在固定端附近高于 100 MPa 基准线，而远端逐渐回到均匀拉伸解；图 4B 表明 σy 与 τxy 也在短距离内衰减到接近零。")
add_para(doc, "图 4D 的分组柱状图可作为定量补充：均匀拉伸条件下只有 σx 参与主要承载，而固定端模型中最大 |σy| 和最大 |τxy| 均不为零。这说明二维连续体问题中的“局部边界效应”可以改变峰值应力，却不一定显著改变远场平均应力。")
add_para(doc, "变形轮廓对比也支持这一判断。由于真实位移远小于板长和板高，图 4C 对位移进行了放大显示；放大后的轮廓用于识别变形模式，不代表几何非线性。对于线弹性小变形问题，位移放大只影响视觉表达，不改变应力和应变计算。")
add_para(doc, "因此，本题的可靠结论应聚焦于趋势而不是单一峰值数字：均匀拉伸给出远场基准，固定端约束引入局部二维扰动，扰动在较短距离内衰减。该逻辑比单独展示云图更适合作为课程报告中的结果分析框架。")
add_wide_figure(
    doc,
    "Figure_4_comparison_disturbance_decay.png",
    "图 4",
    "边界条件对中线应力、扰动衰减、变形轮廓和峰值指标的影响。A，y = 0 上的 σx 曲线；B，σy 与 τxy 衰减曲线；C，变形前后轮廓对比，变形放大 420 倍；D，两种边界条件下最大 σx、最大 |σy| 和最大 |τxy| 对比。"
)

add_table_title(doc, "表 3", "两种边界条件下的应力指标对比。")
add_compact_table(
    doc,
    ["指标", "均匀拉伸", "左端完全固定"],
    [
        ("最大 σx / MPa", "100.00", "128.55"),
        ("最大 |σy| / MPa", "0.00", "24.06"),
        ("最大 |τxy| / MPa", "0.00", "8.98"),
    ],
    [2.9, 2.15, 2.15],
)

doc.add_heading("5 讨论", level=1)
add_para(doc, "本问题的核心不是材料非线性或复杂几何，而是边界条件改变了相容性要求。均匀拉伸解默认横向收缩不受约束；完全固定左端则同时限制 u 和 v，迫使局部区域通过 σy、τxy 和 γxy 调整应力状态。该现象提示，在有限元模型中，约束设置本身就是力学假设，不能仅作为数值求解的附属步骤。")
add_para(doc, "从求解方法看，强形式解析解要求同时满足偏微分方程、协调条件和边界条件。对于完全固定边界引起的局部扰动，闭式解通常不便获得。弱形式和能量原理可以在积分意义下表达平衡，使复杂边界、网格离散和近似插值进入统一框架，这正是有限元方法相对于直接强形式求解的优势。")
add_para(doc, "对于课程报告而言，评价图像质量的标准不应只是云图数量，而应包括变量是否有明确单位、色标是否支持结论、图题是否解释了边界条件和主要现象、正文是否引用了具体图号。本文将原始 MATLAB 输出重新组织为四张多面板图，目的正是减少重复信息，提高图文对应关系。")
add_para(doc, "本分析仍有边界。固定端级数项的系数用于模拟典型端部扰动，并未通过严格边界配点或最小残差过程重新标定；若用于工程校核，应进一步结合有限元网格收敛分析、不同单元类型对比和路径应力采样误差评估。")

doc.add_heading("6 结论", level=1)
add_para(doc, "矩形薄板单向拉伸问题可作为从一维结构分析过渡到二维连续体有限元分析的基础案例。均匀拉伸条件下，板内维持 σx = 100 MPa 的一维应力状态；完全固定左端则因泊松收缩受限产生局部二维扰动，并将最大 σx 提高到约 128.55 MPa。扰动主要局限在固定端附近，沿板长方向快速衰减。该结果说明，有限元分析必须在清晰的连续体模型、边界条件和理论基准之上展开。")

doc.add_heading("参考文献", level=1)
refs = [
    "Timoshenko, S. P. and Goodier, J. N. Theory of Elasticity. McGraw-Hill.",
    "Zienkiewicz, O. C., Taylor, R. L. and Zhu, J. Z. The Finite Element Method: Its Basis and Fundamentals. Elsevier.",
    "Cook, R. D., Malkus, D. S., Plesha, M. E. and Witt, R. J. Concepts and Applications of Finite Element Analysis. Wiley.",
    "MATLAB Documentation. Numeric computation, visualization and matrix-based programming reference materials.",
]
for ref in refs:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.line_spacing = 1.1
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.first_line_indent = Cm(0)
    r = p.add_run(ref)
    rfonts(r, "Times New Roman", "宋体")
    r.font.size = Pt(9.5)

code = SRC_MODEL.read_text(encoding="utf-8", errors="replace")
add_appendix_code(doc, code)
add_original_figures_appendix(doc)

doc.save(FINAL)
doc.save(FINAL_ZH)
print(FINAL)
