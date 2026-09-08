# -*- coding: utf-8 -*-
import json
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


OUT_DIR = r"C:\Users\86198\Desktop\task9"
SUMMARY_PATH = os.path.join(OUT_DIR, "results_summary.json")
DOCX_PATH = os.path.join(OUT_DIR, "有限元分析课程任务9_矩形方板拉伸问题报告.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_width(table, width_dxa=9360, indent_dxa=120):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(width_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_run_font(run, name="Calibri", size=None, bold=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.10):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    for run in paragraph.runs:
        set_run_font(run, size=11)


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    if level == 1:
        before, after, size, color = 16, 8, 16, "2E74B5"
    elif level == 2:
        before, after, size, color = 12, 6, 13, "2E74B5"
    else:
        before, after, size, color = 8, 4, 12, "1F4D78"
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=True, color=color)
    return p


def add_body(doc, text):
    p = doc.add_paragraph()
    p.add_run(text)
    style_paragraph(p)
    return p


def add_bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.line_spacing = 1.167
        p.add_run(text)
        for run in p.runs:
            set_run_font(run, size=11)


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table)
    header_cells = table.rows[0].cells
    for i, header in enumerate(headers):
        header_cells[i].text = header
        set_cell_shading(header_cells[i], "F2F4F7")
        header_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            set_cell_width(header_cells[i], widths[i])
        for p in header_cells[i].paragraphs:
            p.paragraph_format.space_after = Pt(0)
            for run in p.runs:
                set_run_font(run, size=10, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                set_cell_width(cells[i], widths[i])
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=10)
    return table


def fmt_sci(value, unit=""):
    if abs(value) >= 1.0e5 or (abs(value) > 0 and abs(value) < 1.0e-3):
        text = "{:.6e}".format(value)
    else:
        text = "{:.6f}".format(value)
    return text + ((" " + unit) if unit else "")


def add_picture(doc, file_name, caption, width=6.2):
    path = os.path.join(OUT_DIR, file_name)
    if not os.path.exists(path):
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=Inches(width))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(4)
    cap.paragraph_format.space_after = Pt(8)
    r = cap.add_run(caption)
    set_run_font(r, size=10, color="666666")


def main():
    with open(SUMMARY_PATH, "r", encoding="utf-8") as f:
        summary = json.load(f)

    p = summary["parameters"]
    mesh = summary["mesh"]
    theory = summary["theory"]
    result = summary["abaqus_results"]
    load = summary["load"]

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("有限元分析课程任务 9：矩形方板拉伸问题")
    set_run_font(r, size=20, bold=True, color="0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(14)
    r = subtitle.add_run("Abaqus 二维平面应力有限元建模、结果输出与理论对比")
    set_run_font(r, size=11, color="666666")

    add_heading(doc, "一、工程问题与建模参数", 1)
    add_body(doc, "矩形薄板区域为 0 <= x <= L, 0 <= y <= H。板厚远小于平面尺寸，按二维平面应力问题处理；材料为线弹性各向同性材料，静力分析中忽略体力。")
    add_table(
        doc,
        ["项目", "取值"],
        [
            ["长度 L", fmt_sci(p["L_m"], "m")],
            ["高度 H", fmt_sci(p["H_m"], "m")],
            ["厚度 t", fmt_sci(p["t_m"], "m")],
            ["弹性模量 E", fmt_sci(p["E_Pa"], "Pa")],
            ["泊松比 nu", "{:.2f}".format(p["nu"])],
            ["右端均布拉力 q", fmt_sci(p["q_Pa"], "Pa")],
            ["网格尺寸", fmt_sci(p["mesh_size_m"], "m")],
            ["单元类型", p["element_type"] + " 三节点三角形平面应力单元"],
        ],
        widths=[2200, 7160],
    )

    add_heading(doc, "二、Abaqus 有限元模型", 1)
    add_bullets(doc, [
        "几何模型：二维可变形壳/平面应力矩形板，尺寸为 L x H。",
        "截面定义：Homogeneous Solid Section，厚度 t = 0.01 m，并赋予整个矩形板。",
        "材料模型：Elastic，E = 210e9 Pa，nu = 0.30。",
        "装配：创建 Plate-1 实例，用装配边集 LeftEdge 和 RightEdge 管理边界与载荷。",
        "分析步：Static General 静力步。",
        "左边界：x = 0，施加 U1 = 0、U2 = 0 的完全固定条件。",
        "右边界：x = L，按右边界节点的边长权重施加等效节点力，合力等于 q * t * H。",
        "网格：优先使用 CPS3 三节点三角形平面应力单元，网格尺寸 mesh_size = 0.05 m。",
    ])
    add_table(
        doc,
        ["网格/载荷检查项", "Abaqus 输出"],
        [
            ["节点数", mesh["node_count"]],
            ["单元数", mesh["element_count"]],
            ["右边界节点数", mesh["right_edge_node_count"]],
            ["等效右端总力", fmt_sci(load["equivalent_total_force_N"], "N")],
            ["理论右端总力 q*t*H", fmt_sci(load["expected_total_force_N"], "N")],
        ],
        widths=[3200, 6160],
    )

    add_heading(doc, "三、子任务 6：总体组装与边界条件处理", 1)
    add_body(doc, "Abaqus 在求解过程中自动完成单元刚度矩阵组装，形成总体刚度方程 [K]{d} = {F}。左边界固定条件通过 U1 = 0、U2 = 0 引入；右边界均布拉力通过等效节点力进入总体载荷向量。静力求解后得到全体节点位移，再由位移场计算单元应变和应力。")

    add_heading(doc, "四、子任务 7：矩形方板拉伸有限元建模", 1)
    add_body(doc, "本模型的节点自由度为 U1、U2。网格采用二维平面应力单元 CPS3；该单元为三角形常应变单元，理论上与课程中三角形常应变单元推导相对应。若在其他 Abaqus 版本中 CPS3 设置失败，可改用 CPS4R 或 CPS4，但应在报告中说明本任务理论对应 CPS3。")

    add_heading(doc, "五、子任务 8：位移、应变与应力计算", 1)
    add_body(doc, "Abaqus 首先求解得到节点位移 {d}。对单元而言，位移场代入几何矩阵 B 后可计算应变 epsilon = B * d_e；再由平面应力本构矩阵 D 计算应力 sigma = D * epsilon。对于三节点三角形常应变单元，B 矩阵在单元内部为常量，因此单元内应变和应力为常量；不同单元之间的应力可能不同，云图中的变化主要体现边界约束扰动和离散误差。")

    add_heading(doc, "六、子任务 9：理论值与 Abaqus 结果对比", 1)
    add_body(doc, "理想均匀拉伸解析解为 sigma_x = q, sigma_y = 0, tau_xy = 0。平面应力条件下 epsilon_x = q / E, epsilon_y = -nu * q / E, gamma_xy = 0；位移场为 u(x,y) = q*x/E, v(x,y) = -nu*q*y/E。因此右端理论位移 u_right = q*L/E。")
    add_table(
        doc,
        ["对比项目", "理论值", "Abaqus 结果", "相对误差/说明"],
        [
            ["右端平均 U1", fmt_sci(theory["u_right_m"], "m"), fmt_sci(result["U1_right_avg_m"], "m"), "{:.3f}%".format(result["U1_right_error_percent"])],
            ["最大 U1", "-", fmt_sci(result["U1_max_m"], "m"), "出现在右端附近"],
            ["板中部平均 S11", fmt_sci(theory["sigma_x_Pa"], "Pa"), fmt_sci(result["S11_mid_avg_Pa"], "Pa"), "{:.3f}%".format(result["S11_mid_error_percent"])],
            ["S22", "0 Pa", "固定端附近平均绝对值 " + fmt_sci(result["left_region_abs_avg_S22_Pa"], "Pa"), "固定端限制泊松收缩导致局部扰动"],
            ["S12", "0 Pa", "固定端附近平均绝对值 " + fmt_sci(result["left_region_abs_avg_S12_Pa"], "Pa"), "固定端附近可能不为 0，远离固定端趋近 0"],
        ],
        widths=[2100, 2100, 2700, 2460],
    )
    add_body(doc, "从结果看，右端平均 U1 = {:.6e} m，与理论值 {:.6e} m 的误差为 {:.3f}%；板中部平均 S11 = {:.6e} Pa，与 q = {:.6e} Pa 的误差为 {:.3f}%。S22 和 S12 理论上应接近 0，但左端完全固定约束了泊松收缩，因此固定端附近出现局部二维应力扰动；远离固定端后，S11 逐渐接近 q，S22 和 S12 逐渐接近 0。".format(
        result["U1_right_avg_m"],
        theory["u_right_m"],
        result["U1_right_error_percent"],
        result["S11_mid_avg_Pa"],
        theory["sigma_x_Pa"],
        result["S11_mid_error_percent"],
    ))

    add_heading(doc, "七、结果图片与输出文件", 1)
    add_body(doc, "Abaqus 脚本自动保存 ODB、CAE、输入文件、节点位移表、单元应力表、汇总 JSON 和下列结果图片。")
    add_picture(doc, "mesh.png", "图 1  CPS3 三角形网格划分图")
    add_picture(doc, "U1.png", "图 2  x 方向位移 U1 云图")
    add_picture(doc, "U2.png", "图 3  y 方向位移 U2 云图")
    add_picture(doc, "S11.png", "图 4  x 方向正应力 S11 云图")
    add_picture(doc, "S22.png", "图 5  y 方向正应力 S22 云图")
    add_picture(doc, "S12.png", "图 6  剪应力 S12 云图")
    add_picture(doc, "deformed_shape.png", "图 7  变形图")

    add_heading(doc, "八、交付文件清单", 1)
    add_table(
        doc,
        ["文件", "用途"],
        [
            ["task9_abaqus_model.py", "Abaqus 建模、提交、后处理和图片导出脚本"],
            ["task9_rect_plate_tension.cae", "Abaqus/CAE 模型文件"],
            ["task9_rect_plate_tension.odb", "Abaqus 结果数据库"],
            ["node_displacements.csv", "节点坐标及 U1、U2 位移结果"],
            ["element_stresses.csv", "单元/积分点 S11、S22、S12 应力结果"],
            ["results_summary.json", "关键理论值、Abaqus 结果和误差汇总"],
            ["mesh.png, U1.png, U2.png, S11.png, S22.png, S12.png, deformed_shape.png", "网格和结果云图"],
        ],
        widths=[4200, 5160],
    )

    doc.save(DOCX_PATH)
    print(DOCX_PATH)


if __name__ == "__main__":
    main()
