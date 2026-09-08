from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

OUT = Path(r"C:\Users\86198\Desktop\task8_nature_output")
IN_DOC = OUT / "Task8_Nature_optimized_report.docx"
OUT_DOC = OUT / "Task8_Nature_optimized_report_with_original_figures.docx"
ORIG = OUT / "original_figures"

doc = Document(IN_DOC)
doc.add_page_break()
doc.add_heading("附录：原始 MATLAB 图 1-17", level=1)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(8)
r = p.add_run(
    "以下 17 张图来自原始课程报告，用于保留原始计算输出和提交痕迹。正文部分已经将其重新整合为 Figure 1-4 的 Nature-style 多面板科研图。"
)
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(10)

for i, img in enumerate(sorted(ORIG.glob("original_figure_*.png")), 1):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(img), width=Cm(14.5))

    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    r1 = cap.add_run(f"原始图 {i}. ")
    r1.bold = True
    r1.font.name = "Arial"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r1.font.size = Pt(9)
    r2 = cap.add_run("原始 MATLAB 输出图，作为课程报告原始结果记录保留。")
    r2.font.name = "Arial"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(9)

doc.save(OUT_DOC)
print(OUT_DOC)
