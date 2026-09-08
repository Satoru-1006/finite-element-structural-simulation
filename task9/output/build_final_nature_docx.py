import csv
import json
import math
import re
import zipfile
from pathlib import Path

import fitz
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\86198\Desktop\task9")
OUT = ROOT / "output"
ASSETS = OUT / "assets"
FINAL_DOCX = OUT / "final_paper_nature_style.docx"
SOURCE_DOCX = next(p for p in ROOT.glob("*.docx") if not p.name.startswith("~$"))


def ensure_dirs():
    OUT.mkdir(exist_ok=True)
    ASSETS.mkdir(exist_ok=True)


def set_run_font(run, size=10.5, bold=False, italic=False, name="Times New Roman", east="宋体", color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, align=None, before=0, after=3, line=1.15, first_line=True):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if first_line:
        pf.first_line_indent = Cm(0.74)
    if align is not None:
        paragraph.alignment = align


def add_text(doc, text, align=None, before=0, after=3, first_line=True):
    p = doc.add_paragraph()
    set_paragraph(p, align=align, before=before, after=after, first_line=first_line)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    style = "Heading 1" if level == 1 else "Heading 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_before = Pt(8 if level == 1 else 5)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=12 if level == 1 else 10.5, bold=True)
    return p


def add_toc(paragraph):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "目录将在 Word 中自动更新。"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_sep)
    run._r.append(text)
    run._r.append(fld_end)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    set_run_font(run, size=9)


def set_page(section, columns=1):
    section.page_height = Cm(29.7)
    section.page_width = Cm(21.0)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    sect_pr = section._sectPr
    cols = sect_pr.xpath("./w:cols")
    if cols:
        cols = cols[0]
    else:
        cols = OxmlElement("w:cols")
        sect_pr.append(cols)
    cols.set(qn("w:num"), str(columns))
    cols.set(qn("w:space"), "510")


def add_equation(doc, text, num):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(f"{text}    （{num}）")
    set_run_font(r, size=10.5, name="Cambria Math", east="宋体")
    return p


def set_cell(cell, text, bold=False, size=9):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(str(text))
    set_run_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc, text, kind="figure"):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6 if kind == "figure" else 3)
    r = p.add_run(text)
    set_run_font(r, size=9, bold=False)


def add_table(doc, title, headers, rows, widths=None):
    add_caption(doc, title, kind="table")
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        set_cell(table.rows[0].cells[i], h, bold=True, size=8.5)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell(cells[i], val, size=8.5)
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(widths[i] if widths else int(9000 / len(headers))))
            tc_w.set(qn("w:type"), "dxa")
            tc_pr.append(tc_w)
            cell.margin_top = Pt(1)
            cell.margin_bottom = Pt(1)
    return table


def pdf_to_png(pdf_path, out_name):
    out_path = ASSETS / out_name
    if out_path.exists():
        return out_path
    doc = fitz.open(str(pdf_path))
    page = doc[0]
    pix = page.get_pixmap(matrix=fitz.Matrix(3, 3), alpha=False)
    pix.save(str(out_path))
    return out_path


def add_picture(doc, path, caption, width_cm=8.4):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    add_caption(doc, caption, kind="figure")


def extract_code_from_source_docx():
    src = Document(str(SOURCE_DOCX))
    paras = [p.text for p in src.paragraphs]
    start = None
    end = None
    for i, text in enumerate(paras):
        if "%% ============================================================" in text and start is None:
            start = i
        if start is not None and "disp('已导出：task9_element_stress_strain.csv')" in text:
            end = i + 1
            break
    if start is None:
        return Path(r"C:\Users\86198\Desktop\task9\task9_abaqus_model.py").read_text(encoding="utf-8", errors="ignore")
    code = "\n".join(paras[start:end])
    return code.strip()


def read_source_comparison_table():
    src = Document(str(SOURCE_DOCX))
    if len(src.tables) < 2:
        return []
    rows = []
    for r in src.tables[1].rows[1:]:
        rows.append([c.text.strip().replace("\n", " / ") for c in r.cells])
    return rows


def collect_files():
    files = list(ROOT.rglob("*"))
    matlab = [p for p in files if p.suffix.lower() in [".m", ".mlx"]]
    csvs = [p for p in files if p.suffix.lower() in [".csv", ".xlsx"]]
    abaqus = [p for p in files if p.suffix.lower() in [".odb", ".cae", ".inp", ".dat", ".msg", ".sta", ".com", ".prt", ".log", ".rpy", ".rec"]]
    images = [p for p in files if p.suffix.lower() in [".png", ".jpg", ".jpeg", ".pdf"]]
    return matlab, csvs, abaqus, images


def main():
    ensure_dirs()
    summary = json.loads((ROOT / "results_summary.json").read_text(encoding="utf-8"))
    params = summary["parameters"]
    mesh = summary["mesh"]
    theory = summary["theory"]
    abaqus = summary["abaqus_results"]
    load = summary["load"]

    # Convert original MATLAB/PDF figures to PNG assets.
    pdf_figs = {
        "mesh_pdf": (ROOT / "photo" / "三角形网格划分图.pdf", "fig_mesh_matlab.png"),
        "deform_pdf": (ROOT / "photo" / "变形前后对比图.pdf", "fig_deformed_matlab.png"),
        "ux_pdf": (ROOT / "photo" / "x方向位移场 U_x.pdf", "fig_ux_matlab.png"),
        "uy_pdf": (ROOT / "photo" / "y方向位移场 U_y.pdf", "fig_uy_matlab.png"),
        "bc_pdf": (ROOT / "photo'" / "figure2.pdf", "fig_bc.png"),
        "s11_line_pdf": (ROOT / "photo'" / "沿中线 sigma_x 变化曲线.pdf", "fig_s11_line.png"),
        "s11_pdf": (ROOT / "photo'" / "sigma_x 平滑应力云图.pdf", "fig_s11_smooth.png"),
        "s22_pdf": (ROOT / "photo'" / "sigma_y 平滑应力云图.pdf", "fig_s22_smooth.png"),
        "s12_pdf": (ROOT / "photo'" / "tau_{xy} 平滑剪应力云图.pdf", "fig_s12_smooth.png"),
        "s11_const_pdf": (ROOT / "photo'" / "sigma_x 单元常值图.pdf", "fig_s11_const.png"),
    }
    figs = {k: pdf_to_png(path, out) for k, (path, out) in pdf_figs.items() if path.exists()}

    doc = Document()
    set_page(doc.sections[0], columns=1)
    add_page_number(doc.sections[0])

    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    styles["Normal"].font.size = Pt(10.5)
    for style_name in ["Heading 1", "Heading 2"]:
        styles[style_name].font.name = "Times New Roman"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        styles[style_name].font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    r = title.add_run("基于三角形常应变单元的矩形方板拉伸有限元分析")
    set_run_font(r, size=16, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(8)
    r = meta.add_run("有限元分析课程任务 9")
    set_run_font(r, size=10)

    add_heading(doc, "摘要", 1)
    add_text(
        doc,
        "矩形方板拉伸是验证平面弹性有限元方法的典型算例。本文以左边界固定、右边界受 x 方向均布拉应力、上下边界自由的矩形薄板为对象，在平面应力假设下建立三角形常应变单元有限元模型。基于虚功原理和最小势能原理，推导位移插值、形函数矩阵、应变-位移矩阵、平面应力本构矩阵、单元刚度矩阵和右边界等效节点载荷向量，并利用 MATLAB 程序完成总体刚度矩阵组装、位移边界条件处理、节点位移求解及单元应力后处理。结果表明，x 方向位移沿板长方向逐渐增大，y 方向位移体现泊松收缩及固定端约束影响；板中部 σx 与 100 MPa 理论拉应力接近，而 σy 和 τxy 在固定端附近出现局部扰动。该差异主要来源于常应变三角形单元的低阶近似、网格粗细以及左边界完全固定对横向变形的限制。网格加密、局部细化和采用高阶单元可进一步改善位移与应力精度。",
        first_line=False,
    )
    add_text(doc, "关键词：有限元法；三角形常应变单元；平面应力；矩形方板；MATLAB；应力分析", first_line=False)

    add_heading(doc, "目录", 1)
    p = doc.add_paragraph()
    add_toc(p)
    doc.add_page_break()

    sec = doc.add_section(WD_SECTION.CONTINUOUS)
    set_page(sec, columns=2)
    sec.footer.is_linked_to_previous = True

    add_heading(doc, "1 引言", 1)
    add_text(doc, "薄板在拉伸载荷下的变形和应力分布是机械结构设计中的基础问题。对于几何简单、边界条件理想的单向拉伸构件，可由弹性力学给出均匀应力状态下的理论解；然而，当边界约束限制泊松收缩、载荷通过离散节点等效引入或需要观察局部应力扰动时，解析表达通常难以完整描述实际响应。有限元方法通过将连续体区域离散为有限数量的单元，将偏微分方程边值问题转化为代数方程组，因而适合处理复杂边界和局部非均匀应力场。")
    add_text(doc, "本文以矩形方板拉伸问题为算例，采用三节点三角形常应变单元进行平面应力有限元分析。研究内容包括理论弱形式建立、CST 单元公式推导、等效节点载荷构造、总体刚度矩阵组装、MATLAB 数值实现以及与理论均匀拉伸解的对比。文章保留原始课程报告中的真实 MATLAB 程序、网格图、位移场、应力场和结果对比表，并将作业式问答整理为连续论文叙述。")

    add_heading(doc, "2 理论基础", 1)
    add_heading(doc, "2.1 平面弹性问题基本假设", 2)
    add_text(doc, "矩形板长度为 L，高度为 H，厚度为 t，材料为线弹性各向同性材料。由于板厚小于面内尺寸且外载荷作用于板面内，本文采用平面应力假设，即厚度方向应力可忽略。体力项在模型中不计，外载荷主要由右边界沿 x 方向的均布拉应力 q 等效为节点力引入。")
    add_heading(doc, "2.2 虚功原理与最小势能原理", 2)
    add_text(doc, "平面弹性问题的有限元控制方程可由虚功原理建立。对于任意满足位移边界条件的虚位移，结构内部应力虚功应等于体力和边界面力产生的外力虚功。")
    add_equation(doc, "∫Ω δεᵀσ dΩ = ∫Ω δuᵀb dΩ + ∫Γt δuᵀt̄ dΓ", 1)
    add_text(doc, "与虚功原理等价，线弹性稳定结构的真实位移场使总势能在所有许可位移场中取驻值并达到极小。")
    add_equation(doc, "Π(u) = 1/2 ∫Ω εᵀDε dΩ − ∫Ω uᵀb dΩ − ∫Γt uᵀt̄ dΓ", 2)
    add_text(doc, "弱形式降低了对位移试函数光滑性的要求。三角形常应变单元仅保证相邻单元之间位移连续，不保证位移导数连续，因此从虚功或能量形式出发比直接求解强形式更适合。")
    add_heading(doc, "2.3 平面应力本构关系", 2)
    add_text(doc, "在线弹性和小变形条件下，应力与应变满足平面应力本构关系。")
    add_equation(doc, "σ = Dε,  D = E/(1−ν²) [[1,ν,0],[ν,1,0],[0,0,(1−ν)/2]]", 3)

    add_heading(doc, "3 三角形常应变单元理论", 1)
    add_heading(doc, "3.1 节点自由度与位移插值", 2)
    add_text(doc, "三节点三角形单元的节点为 i、j、m，每个节点含有 x 和 y 两个方向的位移自由度。单元节点位移向量写为 de = [ui, vi, uj, vj, um, vm]ᵀ。单元内位移采用一次多项式近似。")
    add_equation(doc, "u(x,y)=a1+a2x+a3y,  v(x,y)=a4+a5x+a6y", 4)
    add_heading(doc, "3.2 形函数矩阵", 2)
    add_text(doc, "由三个节点位移可唯一确定一次位移场，因此单元内任意点位移可由节点位移通过形函数插值得到。")
    add_equation(doc, "{u,v}ᵀ = N de,  N = [[Ni,0,Nj,0,Nm,0],[0,Ni,0,Nj,0,Nm]]", 5)
    add_text(doc, "三角形线性形函数满足节点插值性质和单位分解性质，即在自身节点取 1、其他节点取 0，且 Ni+Nj+Nm=1。")
    add_heading(doc, "3.3 应变-位移矩阵", 2)
    add_text(doc, "平面小变形应变由位移导数给出。由于位移函数为一次多项式，形函数导数在单元内为常数。")
    add_equation(doc, "ε = [εx, εy, γxy]ᵀ = [∂u/∂x, ∂v/∂y, ∂u/∂y+∂v/∂x]ᵀ = Bde", 6)
    add_equation(doc, "B = 1/(2A) [[bi,0,bj,0,bm,0],[0,ci,0,cj,0,cm],[ci,bi,cj,bj,cm,bm]]", 7)
    add_text(doc, "式（7）中，A 为三角形面积，bi、bj、bm 和 ci、cj、cm 由节点坐标差确定。B 矩阵为常量，因此该单元内部应变和应力均为常值，这也是常应变单元名称的来源。")
    add_heading(doc, "3.4 单元刚度矩阵", 2)
    add_text(doc, "将单元位移插值代入虚功方程并整理，可得到单元节点力与节点位移之间的关系。")
    add_equation(doc, "Ke = ∫Ωe BᵀDB t dΩ = t A BᵀDB", 8)
    add_heading(doc, "3.5 等效节点载荷向量", 2)
    add_text(doc, "分布载荷不宜简单平均分到节点，而应通过虚功等效原则转换为节点力。对于右边界长度为 Le 的线性边界单元，均布拉应力 q 沿 x 方向作用时，两端节点在 x 自由度上的等效节点力为 q t Le/2。")
    add_equation(doc, "fe = ∫Γe Nᵀ t̄ t dΓ,  fe,x = [q t Le/2, q t Le/2]ᵀ", 9)

    add_heading(doc, "4 矩形方板有限元模型", 1)
    add_heading(doc, "4.1 几何模型与材料参数", 2)
    add_text(doc, "计算模型采用矩形薄板，材料为线弹性各向同性材料。主要参数见表 1。")
    add_table(
        doc,
        "表 1 矩形方板有限元模型参数",
        ["参数", "取值", "说明"],
        [
            ["L", f"{params['L_m']:.3f} m", "板长"],
            ["H", f"{params['H_m']:.3f} m", "板高"],
            ["t", f"{params['t_m']:.3f} m", "厚度"],
            ["E", f"{params['E_Pa']/1e9:.1f} GPa", "弹性模量"],
            ["ν", f"{params['nu']:.2f}", "泊松比"],
            ["q", f"{params['q_Pa']/1e6:.1f} MPa", "右端均布拉应力"],
            ["单元", "CST / CPS3", "三节点三角形平面应力单元"],
        ],
        widths=[1700, 2300, 3600],
    )
    add_heading(doc, "4.2 边界条件与载荷设置", 2)
    add_text(doc, "左边界施加 u=0、v=0 的完全固定约束，右边界施加沿 x 方向的均布拉力，上下边界保持自由。右边界总等效力为 q t H，数值结果中等效总力为 400000.006 N，与理论总力 400000.000 N 一致。图 1 和图 2 分别给出网格划分以及边界条件与载荷示意。")
    if "mesh_pdf" in figs:
        add_picture(doc, figs["mesh_pdf"], "图 1 矩形方板三角形常应变单元网格划分", width_cm=8.1)
    if "bc_pdf" in figs:
        add_picture(doc, figs["bc_pdf"], "图 2 矩形方板边界条件与载荷示意图", width_cm=8.1)
    add_heading(doc, "4.3 网格划分与自由度编号", 2)
    add_text(doc, "规则矩形区域被划分为三角形单元。每个节点具有两个自由度，若节点编号为 n，则 x 方向和 y 方向自由度编号分别为 2n−1 和 2n。原始报告还比较了 2、4、8 个三角形单元时的结果，用于说明网格粗细对 CST 单元精度的影响。Abaqus 输出网格包含 189 个节点、320 个 CPS3 单元和 9 个右边界加载节点。")
    add_heading(doc, "4.4 总体刚度矩阵组装与边界条件处理", 2)
    add_text(doc, "总体刚度矩阵由各单元刚度矩阵按节点连接关系和自由度编号累加得到。引入边界条件后，将自由自由度和固定自由度分离，在固定自由度位移为零的条件下求解自由自由度位移。")
    add_equation(doc, "Kd = F,  df = Kff⁻¹Ff", 10)

    add_heading(doc, "5 MATLAB 数值实现", 1)
    add_heading(doc, "5.1 程序流程", 2)
    add_text(doc, "MATLAB 程序首先输入几何、材料和载荷参数，并建立平面应力本构矩阵；随后生成规则三角形网格，逐单元计算面积、B 矩阵和 Ke，并组装总体刚度矩阵；最后施加边界条件、求解节点位移、计算单元应变和应力，并输出位移云图、应力云图、变形图和沿中线 σx 变化曲线。程序流程如图 3 所示。")
    # Source Word embedded image8 is the flow chart in the original report.
    source_flow = OUT / "source_media" / "image8.png"
    if source_flow.exists():
        add_picture(doc, source_flow, "图 3 三角形常应变单元有限元分析 MATLAB 程序流程图", width_cm=8.4)
    add_heading(doc, "5.2 节点位移求解", 2)
    add_text(doc, "在获得总体方程后，程序将左边界所有节点的两个位移自由度列入 fixed_dof，并将其余自由度作为 free_dof。节点位移通过求解 K(free_dof, free_dof) d = F(free_dof) 得到。")
    add_heading(doc, "5.3 单元应变与应力计算", 2)
    add_text(doc, "节点位移求得后，单元节点位移向量 de 被代入式（11）计算单元应变，再由式（12）得到单元应力。")
    add_equation(doc, "εe = B de", 11)
    add_equation(doc, "σe = D εe = D B de", 12)
    add_heading(doc, "5.4 后处理与结果输出", 2)
    add_text(doc, "后处理阶段将单元结果平均到节点用于平滑云图显示，同时保留单元常值应力图以说明 CST 单元的原始应力特征。程序输出了节点位移、单元应力、网格图、变形图、位移云图、应力云图和中线应力曲线。")

    add_heading(doc, "6 结果与讨论", 1)
    add_heading(doc, "6.1 网格划分与变形结果", 2)
    add_text(doc, "变形结果见图 4。受右端拉力作用后，矩形板整体沿 x 方向伸长；由于左边界完全固定，靠近固定端的位移被强制为零，变形主要从固定端向加载端逐渐发展。")
    if "deform_pdf" in figs:
        add_picture(doc, figs["deform_pdf"], "图 4 矩形方板变形前后对比图", width_cm=8.1)
    add_heading(doc, "6.2 位移场分析", 2)
    add_text(doc, "如图 5 所示，x 方向位移从左端到右端逐渐增大。这一趋势来自轴向拉伸下的位移累积：左边界位移被约束为零，任意截面的轴向位移可近似理解为从固定端到该截面轴向应变的积分，因此越靠近右端，累计伸长越大。")
    if "ux_pdf" in figs:
        add_picture(doc, figs["ux_pdf"], "图 5 x 方向位移场", width_cm=8.1)
    add_text(doc, "如图 6 所示，y 方向位移反映了泊松效应。板在 x 方向受拉时，材料倾向于在 y 方向收缩；但左边界完全固定同时限制横向位移，使固定端附近的横向变形受到约束，从而形成与理想自由收缩不同的分布。")
    if "uy_pdf" in figs:
        add_picture(doc, figs["uy_pdf"], "图 6 y 方向位移场", width_cm=8.1)
    add_heading(doc, "6.3 应力场分析", 2)
    add_text(doc, "沿中线 σx 变化曲线见图 7。理论均匀拉伸解给出 σx=q=100 MPa，σy=0，τxy=0。有限元结果中，板中部 σx 与理论拉应力整体接近，Abaqus 汇总结果给出的板中部平均 S11 为 100.179 MPa，相对误差约 0.179%。")
    if "s11_line_pdf" in figs:
        add_picture(doc, figs["s11_line_pdf"], "图 7 沿中线 σx 变化曲线", width_cm=8.1)
    add_text(doc, "应力云图见图 8 至图 10。σx 是主要承载应力，远离固定端后趋于均匀拉伸状态；σy 理论上为零，但在固定端附近出现非零值，原因是完全固定边界限制了泊松收缩，使局部区域进入二维应力状态；τxy 理论上也为零，但三角形剖分方向、局部约束扰动和离散化误差会引入小幅剪应力。")
    if "s11_pdf" in figs:
        add_picture(doc, figs["s11_pdf"], "图 8 σx 平滑应力云图", width_cm=8.1)
    if "s22_pdf" in figs:
        add_picture(doc, figs["s22_pdf"], "图 9 σy 平滑应力云图", width_cm=8.1)
    if "s12_pdf" in figs:
        add_picture(doc, figs["s12_pdf"], "图 10 τxy 平滑剪应力云图", width_cm=8.1)
    add_text(doc, "单元常值结果见图 11。由于 CST 单元采用线性位移插值，B 矩阵在单元内为常量，因此单元内应变和应力不会随坐标变化。平滑云图有利于观察整体趋势，但原始单元常值图更直接地反映了 CST 单元的低阶近似特征。")
    if "s11_const_pdf" in figs:
        add_picture(doc, figs["s11_const_pdf"], "图 11 σx 单元常值应力图", width_cm=8.1)
    add_heading(doc, "6.4 理论解与有限元解对比", 2)
    add_text(doc, "原始报告在 y=0 中线上选取 x=0.25 m、0.50 m 和 0.75 m 三个位置进行对比，分别代表靠近固定端影响区、板中部区域和靠近加载端区域，同时避开 x=0 与 x=L 的边界节点。表 2 保留并整理了原始 Word 中的理论解与 2、4、8 个单元有限元结果。")
    add_table(
        doc,
        "表 2 理论解与不同网格有限元结果对比",
        ["比较项目", "理论值", "2 个单元", "4 个单元", "8 个单元", "差异说明"],
        read_source_comparison_table(),
        widths=[1500, 1450, 1450, 1450, 1450, 2350],
    )
    add_text(doc, f"Abaqus 细网格结果进一步表明，右端平均 U1 为 {abaqus['U1_right_avg_m']:.6e} m，理论右端位移为 {theory['u_right_m']:.6e} m，相对误差约 {abaqus['U1_right_error_percent']:.3f}%；板中部平均 S11 为 {abaqus['S11_mid_avg_Pa']/1e6:.3f} MPa，与 100 MPa 理论值接近。")
    add_heading(doc, "6.5 误差来源分析", 2)
    add_text(doc, "主要误差来源见表 3。粗网格下，常应变三角形单元只能以分片常值方式逼近应力场，在固定端和加载端等应力梯度较大的区域误差更明显。左边界完全固定会限制横向泊松收缩，使模型整体刚度偏大，因而位移常表现为偏小，同时诱发固定端附近 σy 和 τxy 的局部扰动。")
    add_table(
        doc,
        "表 3 主要误差来源及影响分析",
        ["误差来源", "主要影响", "改进措施"],
        [
            ["CST 单元低阶近似", "单元内应力为常值，粗网格下应力分布呈块状", "加密网格或采用二次三角形单元"],
            ["固定端约束理想化", "限制泊松收缩，导致位移偏小并产生局部 σy、τxy", "采用更接近实际夹持条件的边界模型"],
            ["边界附近应力梯度", "固定端和加载端附近局部误差较大", "在边界与应力变化明显区域局部细化"],
            ["载荷离散化", "均布力转化为节点力后存在离散近似", "采用更细边界网格并保持虚功等效"],
        ],
        widths=[2300, 3300, 3000],
    )

    add_heading(doc, "7 结论", 1)
    add_text(doc, "本文基于三角形常应变单元建立了矩形方板拉伸问题的平面应力有限元模型，并通过 MATLAB 程序完成了从单元刚度矩阵计算、总体组装、边界条件处理到位移和应力后处理的完整流程。计算结果与理论均匀拉伸解总体一致：x 方向位移沿加载方向逐渐增大，σx 在板中部接近 100 MPa，y 方向位移体现泊松效应。有限元结果中 σy 和 τxy 的局部非零值主要由左端完全固定、网格离散和 CST 单元常应变近似共同造成。对于该类问题，网格加密、边界区域局部细化以及使用高阶单元能够提高位移和应力精度；同时，在报告中应区分理想解析解与实际有限元边界条件之间的适用范围差异。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "[1] Zienkiewicz O C, Taylor R L, Zhu J Z. The Finite Element Method: Its Basis and Fundamentals. 7th ed. Oxford: Butterworth-Heinemann, 2013.",
        "[2] Bathe K J. Finite Element Procedures. Englewood Cliffs: Prentice Hall, 1996.",
        "[3] Cook R D, Malkus D S, Plesha M E, Witt R J. Concepts and Applications of Finite Element Analysis. 4th ed. New York: Wiley, 2001.",
        "[4] Timoshenko S P, Goodier J N. Theory of Elasticity. 3rd ed. New York: McGraw-Hill, 1970.",
        "[5] MATLAB Documentation. MATLAB: The Language of Technical Computing. MathWorks, accessed for numerical linear algebra and visualization usage.",
    ]
    for ref in refs:
        add_text(doc, ref, first_line=False)

    add_heading(doc, "未发现的文件或未采用的资料", 1)
    matlab_files, csv_files, abaqus_files, image_files = collect_files()
    missing = []
    if not matlab_files:
        missing.append("未在源文件夹中发现独立 MATLAB .m 或 .mlx 文件；MATLAB 核心代码来自原始 Word 正文并已置于附录 A。")
    if not any(p.suffix.lower() in [".jpg", ".jpeg"] for p in image_files):
        missing.append("未发现 JPG/JPEG 图片；本文采用源 Word 内嵌 PNG、MATLAB 导出的 PDF 图和 Abaqus 导出的 PNG 图。")
    missing.append("未直接嵌入 Abaqus ODB/CAE 二进制文件；其结果通过已导出的 CSV、JSON 和 PNG 图使用。")
    for item in missing:
        add_text(doc, item, first_line=False)

    add_heading(doc, "附录 A MATLAB 核心代码", 1)
    code = extract_code_from_source_docx()
    for line in code.splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(line)
        set_run_font(r, size=7.5, name="Consolas", east="宋体")

    doc.save(FINAL_DOCX)
    print(FINAL_DOCX)


if __name__ == "__main__":
    main()
