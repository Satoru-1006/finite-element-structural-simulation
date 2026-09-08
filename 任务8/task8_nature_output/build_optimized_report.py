from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUT = Path(r"C:\Users\86198\Desktop\task8_nature_output")
FIG = OUT / "figures"
DOCX = OUT / "任务8_矩形薄板拉伸平面弹性分析_Nature优化版.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_para(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        r1.bold = True
        r1.font.name = "Arial"
        r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r2 = p.add_run(text[len(bold_prefix):])
        r2.font.name = "Arial"
        r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    else:
        r = p.add_run(text)
        r.font.name = "Arial"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    return p


def add_caption(doc, title, caption):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.08
    r1 = p.add_run(title)
    r1.bold = True
    r1.font.name = "Arial"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r1.font.size = Pt(9)
    r2 = p.add_run(" " + caption)
    r2.font.name = "Arial"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r2.font.size = Pt(9)


def add_figure(doc, filename, title, caption, width_cm=16.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(FIG / filename), width=Cm(width_cm))
    add_caption(doc, title, caption)


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(item)
        run.font.name = "Arial"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10)


doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.PORTRAIT
section.page_width = Cm(21.0)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(2.0)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.2)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
styles["Normal"].font.size = Pt(10.5)
for name, size in [("Heading 1", 14), ("Heading 2", 12), ("Heading 3", 10.5)]:
    styles[name].font.name = "Arial"
    styles[name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    styles[name].font.size = Pt(size)
    styles[name].font.bold = True
    styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("矩形薄板单向拉伸问题的平面弹性建模、级数近似与端部扰动分析")
r.bold = True
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(17)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run("有限元分析课程任务 8 优化报告")
r.font.name = "Arial"
r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
r.font.size = Pt(11)
r.italic = True

doc.add_heading("摘要", level=1)
add_para(
    doc,
    "本报告以矩形薄板在右端均布拉力作用下的平面应力问题为对象，重构了从工程问题、连续体力学模型到有限元前处理逻辑的完整分析链条。首先建立二维平面弹性问题的平衡方程、几何方程、本构关系和边界条件；随后利用艾力应力函数说明均匀拉伸解析解的成立条件，并用有限项级数近似描述左端整条边完全固定时产生的端部二维扰动。结果表明，理想均匀拉伸状态下 σx 保持 100 MPa，σy 与 τxy 接近零，位移场分别表现为沿 x 方向的线性伸长和沿 y 方向的泊松收缩。相比之下，完全固定左端抑制了自由泊松收缩，使固定端附近出现局部 σy 与 τxy，并使最大 σx 升至约 128.55 MPa；该扰动随 x 增大迅速衰减，远离固定端后逐渐回到一维均匀拉伸基准。该分析为后续采用有限元弱形式与数值离散求解提供了清晰的理论参照。"
)

doc.add_heading("1. 问题定义与建模目标", level=1)
add_para(
    doc,
    "研究对象为长度 L = 1.0 m、高度 H = 0.4 m、厚度 t = 0.01 m 的矩形薄板。材料取各向同性线弹性体，弹性模量 E = 210 GPa，泊松比 ν = 0.30。板右端承受沿 x 方向的均布拉应力 q = 100 MPa，上下边界自由，体力忽略。由于厚度远小于板面尺寸，厚度方向应力相对于面内应力可忽略，因此本问题按平面应力模型处理。"
)
add_para(
    doc,
    "课程报告的核心目标不是直接给出单个数值答案，而是明确从工程描述到二维弹性边值问题的转换过程。与杆、梁问题相比，平面弹性问题的未知量由单一轴向位移或挠度扩展为二维位移场 u(x,y)、v(x,y)，应变和应力也成为空间连续场变量，因此其控制模型通常表现为偏微分方程组。"
)

add_figure(
    doc,
    "Figure_1_problem_setup.png",
    "Figure 1. Problem setup and analytical framework.",
    "A, 矩形薄板几何、坐标系和主要尺寸。B, 左端完全固定、右端均布拉力和上下自由边界的边界条件。C, 平面应力假设下的面内应力分量示意，其中厚度方向应力近似忽略。D, 从工程问题到平面弹性控制方程、艾力应力函数和位移、应变、应力场输出的求解流程。"
)

doc.add_heading("2. 平面弹性控制方程", level=1)
add_para(
    doc,
    "忽略体力时，二维平面弹性问题的强形式平衡方程为 ∂σx/∂x + ∂τxy/∂y = 0 和 ∂τxy/∂x + ∂σy/∂y = 0。小变形几何关系写作 εx = ∂u/∂x，εy = ∂v/∂y，γxy = ∂u/∂y + ∂v/∂x。对于各向同性平面应力状态，本构关系为 εx = (σx − νσy)/E，εy = (σy − νσx)/E，γxy = τxy/G，其中 G = E/[2(1 + ν)]。"
)
add_para(
    doc,
    "边界条件由位移边界和力边界共同构成。完全固定左端时，x = 0 上满足 u = 0 且 v = 0；右端 x = L 上满足 σx = q、τxy = 0；上下自由边界 y = ±H/2 上满足 σy = 0、τxy = 0。若仅讨论理想均匀拉伸解析解，左端只需去除刚体平移而不约束横向泊松收缩，此时均匀单向应力状态能够与自由收缩相容。"
)

doc.add_heading("3. 艾力应力函数与两类边界解", level=1)
add_para(
    doc,
    "艾力应力函数 Φ 的作用是用单个标量函数生成平面应力分量：σx = ∂²Φ/∂y²，σy = ∂²Φ/∂x²，τxy = −∂²Φ/∂x∂y。该表示可自动满足无体力平衡方程，但为了保证应变协调和位移场存在，Φ 还必须满足双调和方程 ∇⁴Φ = 0。"
)
add_para(
    doc,
    "在均匀拉伸边界下，可取 Φ0 = (q/2)y²，从而得到 σx = q、σy = 0、τxy = 0。对应的平面应力应变为 εx = q/E，εy = −νq/E，γxy = 0；积分后得到位移场 u = qx/E，v = −νqy/E。该解清晰表明，纵向伸长与 x 成正比，而横向位移反映泊松收缩。"
)
add_para(
    doc,
    "当左端整条边完全固定时，均匀拉伸解中的 v = −νqy/E 与 x = 0 上的 v = 0 不相容。为描述这一不相容引起的局部扰动，可将应力场写为均匀拉伸基准项与指数衰减扰动项之和。本文采用 8 项级数近似，扰动项包含 exp(−nπx/H) 和沿高度方向的三角函数项，因此能够表达固定端附近的二维应力重分布，并在远离固定端时快速衰减。该级数解用于揭示机制和趋势，不替代严格解析闭式解。"
)

doc.add_heading("4. 结果", level=1)
doc.add_heading("4.1 均匀拉伸解", level=2)
add_para(
    doc,
    "均匀拉伸解给出了后续比较的理论基准。u 场沿 x 方向线性增加，最大值约为 qL/E = 4.76 × 10⁻⁴ m；v 场仅随 y 改变，表征上下边界向中线方向的泊松收缩。应力场中 σx 在全域内保持 100 MPa，σy 和 τxy 为零；相应的 εx 为 4.76 × 10⁻⁴，εy 为 −1.43 × 10⁻⁴。"
)
add_figure(
    doc,
    "Figure_2_uniform_tension_solution.png",
    "Figure 2. Uniform tension solution.",
    "A, 均匀拉伸条件下 x 方向位移场 u，显示沿板长方向的线性增长。B, y 方向位移场 v，反映由泊松效应导致的横向收缩。C, σx 应力场保持约 100 MPa。D, εx 与 εy 的应变场对照，说明纵向拉伸应变和横向压缩应变在均匀解中均为空间常量。"
)

doc.add_heading("4.2 完全固定左端引起的二维扰动", level=2)
add_para(
    doc,
    "完全固定左端后，板的整体响应不再是纯单向应力状态。由于固定端抑制了横向自由收缩，靠近 x = 0 的区域出现明显二维扰动：σx 在局部区域发生偏离，σy 与 τxy 不再为零，并伴随剪应变 γxy 的局部集中。该扰动主要分布在左端附近，沿 x 方向快速衰减，远端区域逐渐回到均匀拉伸状态。"
)
add_figure(
    doc,
    "Figure_3_fixed_boundary_solution.png",
    "Figure 3. Fully fixed left boundary solution.",
    "A and B, 左端完全固定条件下 u 和 v 的位移场。C, σx 在固定端附近偏离 100 MPa 基准。D and E, σy 与 τxy 使用以零为中心的发散色标，突出正负扰动。F, γxy 剪应变场显示固定端附近的局部剪切响应。灰色阴影标出 fixed-end disturbance zone。"
)

doc.add_heading("4.3 两种边界条件的定量比较", level=2)
add_para(
    doc,
    "沿中线 y = 0 的曲线显示，均匀拉伸解析解在全长范围内保持 σx = 100 MPa；完全固定左端的级数近似解则在固定端附近产生更高的 σx，随后迅速衰减到 100 MPa 基准线。σy 与 τxy 的中线曲线也具有相同的端部局部化特征，其幅值随 x 增大快速接近零。"
)
add_para(
    doc,
    "本模型下，固定端边界条件使最大 σx 约为 128.55 MPa，最大 |σy| 约为 24.06 MPa，最大 |τxy| 约为 8.98 MPa。均匀拉伸解则只有 σx = 100 MPa，σy 与 τxy 为零。该差异说明，边界条件不仅影响位移约束，还会改变应力路径和局部峰值，这也是有限元分析中必须谨慎处理约束方式的原因。"
)
add_figure(
    doc,
    "Figure_4_comparison_disturbance_decay.png",
    "Figure 4. Comparison and disturbance decay.",
    "A, 沿中线 y = 0 的 σx 曲线对比，黑色实线为均匀拉伸基准，蓝色虚线为左端完全固定级数近似解，100 MPa 基准线用于标定远场应力。B, σy 与 τxy 沿中线的衰减曲线，显示固定端扰动的局部化。C, 变形前后轮廓对比，变形量按 420 倍放大以便观察。D, 两种边界条件下最大 σx、最大 |σy| 和最大 |τxy| 的分组对比。"
)

doc.add_heading("5. 讨论", level=1)
add_para(
    doc,
    "该结果强调了连续体有限元建模中的一个关键问题：边界条件是力学模型的一部分，而不是后处理细节。理想均匀拉伸解隐含板可以自由发生泊松收缩，因此能够保持一维应力状态；完全固定左端则额外施加了横向位移约束，破坏了这一相容性，迫使局部区域通过 σy、τxy 和 γxy 来重新满足平衡与协调。"
)
add_para(
    doc,
    "从有限元角度看，强形式解析解要求同时满足偏微分方程和全部边界条件，实际工程边界往往很难得到闭式表达。弱形式或能量原理将控制方程转化为积分意义下的平衡，使复杂区域、复杂边界和离散插值可以统一处理。因此，本题的解析解和级数近似更适合作为有限元计算的理论参照与结果解释框架，而不是替代数值离散。"
)
add_para(
    doc,
    "本报告中的固定端级数解采用有限项近似，系数用于刻画端部扰动的典型衰减规律；若要获得严格工程校核结果，应进一步通过有限元网格收敛分析、边界条件敏感性分析和与解析基准的定量误差比较来验证。"
)

doc.add_heading("6. 结论", level=1)
add_bullets(
    doc,
    [
        "矩形薄板单向拉伸在薄板假设下可归结为平面应力边值问题，基本未知量为 u(x,y) 和 v(x,y)。",
        "均匀拉伸解析解给出 σx = 100 MPa、σy = τxy = 0，并产生线性纵向伸长和横向泊松收缩。",
        "左端完全固定会破坏自由泊松收缩，使固定端附近产生 σy、τxy 和 γxy 局部扰动，扰动随 x 快速衰减。",
        "最大 σx 从均匀解的 100 MPa 增至约 128.55 MPa，说明约束方式会显著影响局部应力峰值。",
        "该问题展示了从连续体强形式到有限元弱形式的必要性：复杂边界条件下，数值离散是获得可靠工程解的主要途径。"
    ]
)

doc.add_heading("符号与参数表", level=1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
hdr = table.rows[0].cells
for cell, text in zip(hdr, ["符号", "数值或表达式", "说明"]):
    set_cell_text(cell, text, bold=True)
    set_cell_shading(cell, "EDEDED")
rows = [
    ("L", "1.0 m", "板长"),
    ("H", "0.4 m", "板高"),
    ("t", "0.01 m", "板厚"),
    ("E", "210 GPa", "弹性模量"),
    ("ν", "0.30", "泊松比"),
    ("q", "100 MPa", "右端均布拉应力"),
    ("G", "E/[2(1 + ν)]", "剪切模量"),
    ("σx, σy, τxy", "MPa", "平面应力分量"),
    ("u, v", "m", "x 与 y 方向位移")
]
for row in rows:
    cells = table.add_row().cells
    for cell, text in zip(cells, row):
        set_cell_text(cell, text)

doc.add_heading("图件与源数据说明", level=1)
add_para(
    doc,
    "本报告所有插图均由 R 生成，PNG 版本按 300 dpi 导出，另提供 PDF 与 SVG 矢量版本。源数据包括全场解析/级数场变量、沿中线应力曲线和应力指标汇总，可用于复核图中数值。"
)

doc.save(DOCX)
print(DOCX)
